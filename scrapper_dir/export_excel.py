import os
from openpyxl import Workbook
from scrapper_dir.db import get_all_data

# Создание файла и активного листа
wb = Workbook()
ws = wb.active
ws.title = 'Выгрузка БД'

# Заполняем название колонок в Excel
ws['A1'] = '№ п/п'
ws['B1'] = 'Название валюты'
ws['C1'] = 'Обозначение валюты'
ws['D1'] = 'Цена - $'
ws['E1'] = 'Объём торгов'
ws['F1'] = 'Капитализация'
ws['G1'] = 'Дата добавления информации'

# Данные для сохранения файла
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "cryptocurrency data.xlsx")


def export_to_excel():
    """Выполняет перенос данных из БД в Excel"""
    coins_data = get_all_data()
    for row in coins_data:
        ws.append(row)
    wb.save(EXCEL_PATH)
    